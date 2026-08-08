from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Fooddely_Essential_API_Requirements.docx"

RED, DARK, MUTED = "C83B2B", "241F1C", "6D625C"
PALE, LINE, WHITE = "FFF0EB", "E4D8D2", "FFFFFF"


def font(run, size=10.2, bold=False, color=DARK, name="Aptos"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def margins(cell, value=120):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.append(mar)


def geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for i, width in enumerate(widths):
            row.cells[i].width = Inches(width)
            tcpr = row.cells[i]._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(int(width * 1440)))
            tcw.set(qn("w:type"), "dxa")
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tblw.set(qn("w:type"), "dxa")


def borders(table, color=LINE):
    tblpr = table._tbl.tblPr
    node = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), color)
        node.append(border)
    tblpr.append(node)


def paragraph(doc, text, after=6, color=DARK):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    font(p.add_run(text), color=color)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(3)
    font(p.add_run(text), size=10)


def code(doc, content):
    table = doc.add_table(rows=1, cols=1)
    geometry(table, [6.5])
    borders(table)
    cell = table.cell(0, 0)
    shade(cell, "F8F6F5")
    margins(cell, 140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    lines = content.strip().splitlines()
    for i, line in enumerate(lines):
        run = p.add_run(line)
        font(run, size=8.1, name="Courier New")
        if i < len(lines) - 1:
            run.add_break()
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    geometry(table, [6.5])
    cell = table.cell(0, 0)
    shade(cell, PALE)
    margins(cell, 150)
    p = cell.paragraphs[0]
    font(p.add_run(f"{label}: "), size=10, bold=True, color=RED)
    font(p.add_run(text), size=10)


def api(doc, n, title, reason, endpoint, request=None, response=None, rules=None):
    heading = doc.add_heading(f"{n}. {title}", level=1)
    heading.paragraph_format.keep_with_next = True
    tag = doc.add_paragraph()
    tag.paragraph_format.space_after = Pt(5)
    font(tag.add_run("ESSENTIAL"), size=8.3, bold=True, color=RED)
    font(tag.add_run("  •  Required for complete customer flow"), size=8.3, color=MUTED)
    paragraph(doc, reason)
    doc.add_heading("Endpoint", level=2)
    code(doc, endpoint)
    if request:
        doc.add_heading("Request", level=2)
        code(doc, request)
    if response:
        doc.add_heading("Required response", level=2)
        code(doc, response)
    if rules:
        doc.add_heading("Required behavior", level=2)
        for rule in rules:
            bullet(doc, rule)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.4)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.2)
normal.font.color.rgb = RGBColor.from_string(DARK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for style_name, size, before, after, color in (
    ("Heading 1", 15.5, 15, 6, RED),
    ("Heading 2", 11.2, 8, 4, DARK),
    ("Heading 3", 10.5, 7, 3, DARK),
):
    style = doc.styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

hp = section.header.paragraphs[0]
font(hp.add_run("FOOD"), size=8.5, bold=True, color=RED)
font(hp.add_run("DELY  |  Essential API Requirements"), size=8.5, bold=True)

fp = section.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(fp.add_run("Backend handoff  •  "), size=8, color=MUTED)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
fp._p.append(field)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(22)
p.paragraph_format.space_after = Pt(4)
font(p.add_run("BACKEND IMPLEMENTATION HANDOFF"), size=9, bold=True, color=RED)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
font(p.add_run("Fooddely Essential API Work"), size=27, bold=True, name="Aptos Display")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
font(
    p.add_run("Only the backend changes required to complete the customer journey"),
    size=13.5,
    color=MUTED,
    name="Aptos Display",
)

for label, value in (
    ("Scope", "Authentication, addresses, restaurants, menu, cart, discounts, Stripe and order tracking"),
    ("Audience", "Backend developer / project client"),
    ("Prepared", date.today().strftime("%d %B %Y")),
    ("Priority", "All items in this document are required"),
):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    font(p.add_run(f"{label}: "), bold=True)
    font(p.add_run(value), color=MUTED)

doc.add_paragraph()
callout(
    doc,
    "Definition of complete",
    "A customer can register or sign in, manage an address, find a restaurant, configure products, maintain a correctly priced cart, apply discounts, pay through Stripe, receive exactly one order and track its live status.",
)

doc.add_heading("Required work at a glance", level=1)
rows = [
    ("P0", "Cart quantity and totals", "Blocks reliable cart and secure checkout."),
    ("P0", "Stripe intent, confirm and webhook", "Blocks real payment and order creation."),
    ("P0", "Order status APIs", "Blocks live customer tracking."),
    ("P1", "Login refresh and guest-cart merge", "Prevents session loss and cart loss."),
    ("P1", "Registration/reset/address update", "Completes account management."),
    ("P1", "Menu/options/discount validation", "Ensures server-authoritative ordering."),
]
table = doc.add_table(rows=1, cols=3)
geometry(table, [0.55, 2.25, 3.7])
borders(table)
for i, value in enumerate(("Priority", "Backend work", "Why required")):
    cell = table.rows[0].cells[i]
    shade(cell, DARK)
    margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    font(cell.paragraphs[0].add_run(value), size=9, bold=True, color=WHITE)
for priority, work, why in rows:
    cells = table.add_row().cells
    for cell in cells:
        margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cells[0], PALE)
    font(cells[0].paragraphs[0].add_run(priority), size=9, bold=True, color=RED)
    font(cells[1].paragraphs[0].add_run(work), size=9.1, bold=True)
    font(cells[2].paragraphs[0].add_run(why), size=9.1, color=MUTED)

api(
    doc, 1, "Login identity and token refresh",
    "The login response must identify the account role and support access-token renewal.",
    """POST /api/token/
POST /api/token/refresh/""",
    """// Login
{"username": "customer@example.com", "password": "password"}

// Refresh
{"refresh": "<refresh-token>"}""",
    """{
  "access": "<access-token>",
  "refresh": "<refresh-token>",
  "expires_in": 900,
  "user": {
    "id": 18,
    "email": "customer@example.com",
    "firstname": "John",
    "lastname": "Smith",
    "role": "customer"
  }
}""",
    [
        "Return one explicit role: customer, restaurant_owner or admin.",
        "Return a new access token from the refresh endpoint.",
        "Prefer a secure HTTP-only cookie for the refresh token.",
    ],
)

api(
    doc, 2, "Customer registration and password reset",
    "Registration needs stable field validation, and the current forgot-password screen has no real backend flow.",
    """POST /api/app/customer/register/
POST /api/app/auth/password-reset/
POST /api/app/auth/password-reset/validate/
POST /api/app/auth/password-reset/confirm/""",
    """// Registration
{
  "firstname": "John",
  "lastname": "Smith",
  "birthday": "1995-08-17",
  "address": "12 Bahnhofstrasse",
  "postal_code": "8001",
  "city": "Zürich",
  "email": "john@example.com",
  "phone": "+41791234567",
  "password": "password",
  "password_confirm": "password"
}

// Reset confirmation
{
  "uid": "encoded-user-id",
  "token": "single-use-token",
  "password": "new-password",
  "password_confirm": "new-password"
}""",
    """{
  "code": "validation_error",
  "message": "Some details need attention.",
  "fields": {
    "email": ["An account already exists with this email."],
    "phone": ["Enter a valid international phone number."]
  }
}""",
    [
        "Normalize email and phone values and enforce uniqueness.",
        "Use short-lived, single-use reset tokens.",
        "Return a generic reset-request response to prevent account enumeration.",
    ],
)

api(
    doc, 3, "Address update and default management",
    "Address list/create/delete exist; editing and changing the default address are still required.",
    "PATCH /api/app/customer/addresses/{address_id}/",
    """{
  "address": "12 Bahnhofstrasse",
  "postal_code": "8001",
  "city": "Zürich",
  "default": true
}""",
    """{
  "id": 17,
  "address": "12 Bahnhofstrasse",
  "postal_code": "8001",
  "city": "Zürich",
  "default": true
}""",
    [
        "Verify address ownership.",
        "Unset the previous default in the same transaction.",
        "Protect the last/default address during deletion.",
    ],
)

api(
    doc, 4, "Restaurant listing and filters",
    "The listing endpoint must support the filters already represented by the customer experience.",
    """GET /api/app/restaurants/
  ?search=pizza
  &category=4
  &open_now=true
  &delivery_available=true
  &pickup_available=true
  &no_waste=false
  &page=1
  &page_size=20""",
    response="""{
  "count": 83,
  "next": "/api/app/restaurants/?page=2",
  "previous": null,
  "results": [{
    "id": 7,
    "name": "Restaurant",
    "open": true,
    "delivery_available": true,
    "pickup_available": true,
    "delivery_fee": "3.50",
    "delivery_time": 35,
    "minimum_order": "20.00",
    "rating": "4.7",
    "reviews": 128,
    "categories": [],
    "images": []
  }]
}""",
)

api(
    doc, 5, "Restaurant menu, options and allergens",
    "Restaurant details must expose complete, consistent product configuration and availability.",
    "GET /api/app/restaurants/{restaurant_id}/",
    response="""{
  "id": 364,
  "name": "Classic Burger",
  "price": "18.00",
  "image": null,
  "available": true,
  "allergies": [{"id": 2, "name": "Gluten"}],
  "options": [{
    "id": 397,
    "name": "Choose cheese",
    "required": true,
    "multiple": false,
    "minimum_selections": 1,
    "maximum_selections": 1,
    "items": [{
      "id": 1190,
      "name": "Gruyère",
      "price": "2.50",
      "available": true
    }]
  }]
}""",
    rules=[
        "Validate required options and selection limits when adding to cart.",
        "Reject unavailable products and option items.",
        "Return an accessible image URL or null; do not return URLs that respond with 403.",
    ],
)

api(
    doc, 6, "Cart quantity mutation",
    "PATCH is currently unsupported. The frontend must temporarily delete and recreate a line, which is not atomic.",
    """PATCH /api/app/cart/items/{cart_item_id}/
Authorization: Bearer <access-token>
X-Session-ID: <guest-session-id>""",
    """{"quantity": 3}""",
    """{
  "id": 42,
  "items": [],
  "item_count": 3,
  "subtotal": "42.50",
  "discount_total": "5.00",
  "delivery_fee": "3.50",
  "total_price": "41.00",
  "currency": "CHF"
}""",
    [
        "Verify line ownership.",
        "Validate quantity, stock, product status and selected options.",
        "Return the complete recalculated cart.",
    ],
)

api(
    doc, 7, "Complete cart pricing and deduplication",
    "Cart responses must never return null totals, and identical add-to-cart requests must not create duplicate lines.",
    """GET /api/app/cart/
POST /api/app/cart/
DELETE /api/app/cart/""",
    response="""{
  "id": 42,
  "items": [{
    "id": 501,
    "quantity": 2,
    "unit_price": "18.00",
    "options_price": "2.50",
    "line_total": "41.00",
    "menu_item": {},
    "options": []
  }],
  "item_count": 2,
  "subtotal": "41.00",
  "discount_total": "0.00",
  "delivery_fee": "3.50",
  "total_price": "44.50",
  "currency": "CHF"
}""",
    rules=[
        "All money fields are decimal strings and never null.",
        "Combine items only when product, restaurant, options and instructions match.",
        "Support an Idempotency-Key header for add-to-cart retries.",
    ],
)

api(
    doc, 8, "Guest cart merge after login",
    "Guest cart contents must survive customer authentication.",
    """POST /api/app/cart/merge/
Authorization: Bearer <customer-token>""",
    """{"session_id": "guest-session-id"}""",
    """{
  "id": 42,
  "customer": 18,
  "session_id": null,
  "items": [],
  "subtotal": "52.00",
  "total_price": "55.50",
  "currency": "CHF"
}""",
    [
        "Combine identical lines.",
        "Return a structured conflict if customer and guest carts belong to different restaurants.",
        "Invalidate the guest session only after a successful merge.",
    ],
)

api(
    doc, 9, "Discount and checkout quote",
    "Discounts, delivery eligibility and final totals must be validated by the backend before Stripe intent creation.",
    "POST /api/app/cart/quote/",
    """{
  "fulfilment_type": "delivery",
  "address_id": 17,
  "postal_code": "8001",
  "preferred_time": "2026-07-30T20:30:00+05:00",
  "discount_code": "WELCOME10"
}""",
    """{
  "valid": true,
  "errors": [],
  "subtotal": "52.00",
  "discount_total": "5.00",
  "delivery_fee": "3.50",
  "total_price": "50.50",
  "currency": "CHF",
  "minimum_order_met": true,
  "restaurant_open": true,
  "items_available": true
}""",
)

api(
    doc, 10, "Stripe PaymentIntent",
    "The server must create the PaymentIntent from the validated cart and persist fulfilment details.",
    """POST /api/app/payment/intent/
Authorization: Bearer <access-token>
X-Session-ID: <guest-session-id>""",
    """{
  "fulfilment_type": "delivery",
  "delivery_firstname": "John",
  "delivery_lastname": "Smith",
  "delivery_address": "12 Bahnhofstrasse",
  "delivery_postal_code": "8001",
  "delivery_city": "Zürich",
  "delivery_phone": "+41791234567",
  "delivery_email": "john@example.com",
  "preferred_time": "2026-07-30T20:30:00+05:00",
  "delivery_notes": "Ring apartment 4",
  "receive_notifications": false
}""",
    """{
  "payment_intent_id": "pi_...",
  "payment_intent_client_secret": "pi_..._secret_...",
  "cart_id": 42,
  "subtotal": "52.00",
  "discount_total": "5.00",
  "delivery_fee": "3.50",
  "total_price": "50.50",
  "currency": "CHF"
}""",
    [
        "Calculate the Stripe amount only on the backend.",
        "Validate cart, availability, minimum order and delivery area first.",
        "For pickup, address fields may be optional while contact fields remain required.",
    ],
)

api(
    doc, 11, "Stripe confirmation and webhook",
    "Payment confirmation must create exactly one order and remain reliable if the browser closes or retries.",
    """POST /api/app/payment/confirm/
POST /api/webhooks/stripe/""",
    """// Confirmation
{"payment_intent_id": "pi_..."}""",
    """{
  "id": 1042,
  "status": "PLACED",
  "payment_intent_id": "pi_...",
  "payment_status": "succeeded",
  "price": "50.50",
  "currency": "CHF",
  "delivery": {},
  "items": []
}""",
    [
        "Verify PaymentIntent amount, currency and status directly with Stripe.",
        "Make confirmation idempotent: one order per PaymentIntent.",
        "Clear or submit the cart atomically with order creation.",
        "Verify Stripe webhook signatures.",
        "Handle payment_intent.succeeded, payment_intent.payment_failed, payment_intent.canceled and charge.refunded.",
    ],
)

api(
    doc, 12, "Customer orders and live tracking",
    "Order history, detail and status polling must return ownership-protected, consistent data.",
    """GET  /api/app/orders/
GET  /api/app/orders/{order_id}/
GET  /api/app/orders/{order_id}/status/
POST /api/app/orders/{order_id}/cancel/""",
    response="""{
  "id": 1042,
  "status": "PREPARING",
  "updated_at": "2026-07-30T20:15:00Z",
  "estimated_delivery_at": "2026-07-30T20:50:00Z"
}""",
    rules=[
        "Verify customer ownership for list, detail, status and cancellation.",
        "Use a stable status set: PLACED, PREPARING, DELIVERED, COMPLETED and CANCELLED.",
        "Return complete item, option, price, delivery and timestamp data in order details.",
        "Only permit cancellation during allowed states and return a structured rejection otherwise.",
    ],
)

doc.add_heading("Required configuration", level=1)
paragraph(doc, "Frontend environment:")
code(doc, """NEXT_PUBLIC_API_URL=https://backend.fooddely.com
NEXT_PUBLIC_STRIPE_PUBLISHABLE_KEY=pk_test_...""")
paragraph(doc, "Backend environment:")
code(doc, """STRIPE_SECRET_KEY=sk_test_...
STRIPE_WEBHOOK_SECRET=whsec_...""")
callout(
    doc,
    "Security",
    "The Stripe secret key and webhook secret must remain on the backend. They must never be committed to the frontend repository or exposed to browser code.",
)

doc.add_heading("Acceptance checklist", level=1)
for item in (
    "Registration, login, token refresh and reset work with structured errors.",
    "Customer addresses can be created, edited, selected as default and deleted safely.",
    "Restaurant search and filters return complete, paginated data.",
    "Menu availability, required options and allergens are enforced server-side.",
    "Cart quantity changes are atomic and every price field is non-null.",
    "Guest cart merges after login without losing items.",
    "Discount quote and final checkout totals come from the backend.",
    "Stripe test payment creates exactly one order and webhook processing is verified.",
    "Customer can view order history and receive live status updates.",
):
    bullet(doc, item)

doc.core_properties.title = "Fooddely Essential API Requirements"
doc.core_properties.subject = "Only required backend APIs for complete customer flow"
doc.core_properties.author = "Fooddely Frontend Team"
doc.save(OUTPUT)
print(OUTPUT)
