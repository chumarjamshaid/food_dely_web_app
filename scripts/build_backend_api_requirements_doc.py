from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Fooddely_Backend_API_Requirements.docx"

RED = "C83B2B"
DARK = "241F1C"
MUTED = "6D625C"
LIGHT = "F7F1EE"
PALE_RED = "FFF0EB"
LINE = "E4D8D2"
GREEN = "1F7A53"
AMBER = "9A6700"
BLUE = "1F4D78"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color=LINE, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)
            tc_pr = row.cells[index]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")


def format_run(run, size=10.5, bold=False, color=DARK, font="Aptos"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, *, bold_prefix=None, after=6, color=DARK, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if bold_prefix and text.startswith(bold_prefix):
        format_run(p.add_run(bold_prefix), size=size, bold=True, color=color)
        format_run(p.add_run(text[len(bold_prefix):]), size=size, color=color)
    else:
        format_run(p.add_run(text), size=size, color=color)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    format_run(p.add_run(text), size=10.25)
    return p


def add_code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_widths(table, [6.5])
    set_table_borders(table, color=LINE, size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8F6F5")
    set_cell_margins(cell, top=110, start=150, bottom=110, end=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for idx, line in enumerate(text.strip().splitlines()):
        run = p.add_run(line)
        format_run(run, size=8.2, color=DARK, font="Courier New")
        if idx < len(text.strip().splitlines()) - 1:
            run.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_callout(doc, label, text, fill=PALE_RED, color=RED):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_widths(table, [6.5])
    set_table_borders(table, color=fill, size="2")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=170, bottom=130, end=170)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    format_run(p.add_run(f"{label}: "), size=10, bold=True, color=color)
    format_run(p.add_run(text), size=10, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_api_section(doc, number, title, priority, purpose, endpoint, request=None,
                    response=None, requirements=None, note=None):
    h = doc.add_heading(f"{number}. {title}", level=1)
    h.paragraph_format.keep_with_next = True
    badge = doc.add_paragraph()
    badge.paragraph_format.space_after = Pt(6)
    priority_color = RED if priority == "REQUIRED" else AMBER
    format_run(badge.add_run(priority), size=8.5, bold=True, color=priority_color)
    format_run(badge.add_run("  •  Backend API contract"), size=8.5, color=MUTED)
    add_text(doc, purpose, after=7)
    sub = doc.add_heading("Endpoint", level=2)
    sub.paragraph_format.keep_with_next = True
    add_code(doc, endpoint)
    if request:
        sub = doc.add_heading("Request", level=2)
        sub.paragraph_format.keep_with_next = True
        add_code(doc, request)
    if response:
        sub = doc.add_heading("Response", level=2)
        sub.paragraph_format.keep_with_next = True
        add_code(doc, response)
    if requirements:
        sub = doc.add_heading("Backend requirements", level=2)
        sub.paragraph_format.keep_with_next = True
        for item in requirements:
            add_bullet(doc, item)
    if note:
        add_callout(doc, "Implementation note", note)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.4)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(DARK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for style_name, size, before, after, color in (
    ("Heading 1", 16, 16, 7, RED),
    ("Heading 2", 11.5, 9, 4, DARK),
    ("Heading 3", 10.5, 7, 3, BLUE),
):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
format_run(hp.add_run("FOOD"), size=9, bold=True, color=RED)
format_run(hp.add_run("DELY  |  Backend API Requirements"), size=9, bold=True, color=DARK)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
format_run(fp.add_run("Confidential technical handoff  •  "), size=8, color=MUTED)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
fp._p.append(field)

# Cover / memo masthead
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(20)
p.paragraph_format.space_after = Pt(5)
format_run(p.add_run("BACKEND DELIVERY SPECIFICATION"), size=9, bold=True, color=RED)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
format_run(p.add_run("Fooddely Backend API"), size=27, bold=True, color=DARK, font="Aptos Display")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
format_run(p.add_run("Required modifications and new endpoints for completed frontend flows"),
           size=14, color=MUTED, font="Aptos Display")

metadata = [
    ("Prepared for", "Backend development team / project client"),
    ("Frontend scope", "Authentication, addresses, restaurants, menus, cart and checkout"),
    ("Prepared on", date.today().strftime("%d %B %Y")),
    ("Document status", "Implementation requirements — ready for backend estimation"),
]
for label, value in metadata:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    format_run(p.add_run(f"{label}: "), size=10, bold=True, color=DARK)
    format_run(p.add_run(value), size=10, color=MUTED)

doc.add_paragraph().paragraph_format.space_after = Pt(4)
add_callout(
    doc,
    "Objective",
    "Remove frontend workarounds, stabilize the complete customer journey, and make the backend the authoritative source for pricing, authentication, availability and payment state.",
)

doc.add_heading("Executive summary", level=1)
add_text(
    doc,
    "The frontend is operational, but several backend gaps currently require non-atomic workarounds or prevent complete flows. The highest-risk gaps are cart quantity updates, missing cart totals, payment contract completeness, guest-cart preservation, token refresh and password reset.",
)
add_text(
    doc,
    "All financial values must be calculated and validated by the backend. The frontend may display temporary estimates, but it must never be trusted as the source of the Stripe amount, discounts, fees, taxes or inventory availability.",
)

doc.add_heading("Delivery priority", level=1)
priority_rows = [
    ("P0", "Cart quantity update", "Current PATCH request is unsupported; update requires delete + re-add."),
    ("P0", "Cart totals and line pricing", "Backend returned null total; secure checkout needs authoritative totals."),
    ("P0", "Payment intent and confirmation", "Needs complete fulfilment fields and idempotent order creation."),
    ("P1", "Guest-cart merge", "Guest items can be lost when the customer signs in."),
    ("P1", "Authentication refresh and reset", "No refresh flow; forgot-password screen has no real API."),
    ("P1", "Address editing/default", "Create/list/delete exist; update/default management is missing."),
    ("P1", "Image URL reliability", "Some menu image URLs return 403."),
    ("P2", "Restaurant/menu consistency", "Normalize availability, option and allergen payloads."),
    ("P2", "Search, filters and quote endpoint", "Improves scalability and checkout validation."),
]
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
set_table_widths(table, [0.55, 1.95, 4.0])
set_table_borders(table)
headers = ("Priority", "Workstream", "Reason")
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    set_cell_shading(cell, DARK)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    format_run(p.add_run(text), size=9, bold=True, color=WHITE)
set_repeat_table_header(table.rows[0])
for priority, workstream, reason in priority_rows:
    cells = table.add_row().cells
    for cell in cells:
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cells[0], PALE_RED if priority == "P0" else ("FFF7E6" if priority == "P1" else LIGHT))
    format_run(cells[0].paragraphs[0].add_run(priority), size=9, bold=True, color=RED if priority == "P0" else DARK)
    format_run(cells[1].paragraphs[0].add_run(workstream), size=9.2, bold=True)
    format_run(cells[2].paragraphs[0].add_run(reason), size=9.2, color=MUTED)

doc.add_heading("Detailed API contracts", level=1)
add_text(doc, "Unless otherwise stated, authenticated requests use Authorization: Bearer <access-token>. Guest cart and order requests use X-Session-ID and may also accept session_id as a query parameter during migration.")

add_api_section(
    doc, 1, "Cart quantity update", "REQUIRED",
    "The cart API currently allows GET, POST and DELETE only. Quantity changes therefore require deleting and recreating a line, which is not atomic and may lose the item if the second request fails.",
    """PATCH /api/app/cart/items/{cart_item_id}/
Authorization: Bearer <access-token>
X-Session-ID: <guest-session-id>
Content-Type: application/json""",
    """{
  "quantity": 3
}""",
    """{
  "id": 42,
  "customer": 18,
  "session_id": null,
  "items": [],
  "item_count": 3,
  "subtotal": "42.50",
  "discount_total": "5.00",
  "delivery_fee": "3.50",
  "total_price": "41.00",
  "currency": "CHF"
}""",
    [
        "quantity is a required integer with a minimum of 1 and a defined maximum such as 99.",
        "Verify that the line belongs to the authenticated customer or guest session.",
        "Revalidate product availability, option availability, stock and restaurant status.",
        "Return the complete recalculated cart after the mutation.",
    ],
)

add_api_section(
    doc, 2, "Authoritative cart totals", "REQUIRED",
    "Guest-cart testing returned total_price as null. The backend must always return complete decimal totals because the frontend cannot securely calculate option prices, discounts, delivery fees, taxes or the Stripe amount.",
    """GET /api/app/cart/
Authorization: Bearer <access-token>
X-Session-ID: <guest-session-id>""",
    response="""{
  "id": 42,
  "customer": null,
  "session_id": "guest-session-id",
  "items": [],
  "item_count": 4,
  "subtotal": "52.00",
  "options_total": "4.50",
  "discount_total": "5.00",
  "delivery_fee": "3.50",
  "service_fee": "0.00",
  "tax_total": "0.00",
  "total_price": "55.00",
  "currency": "CHF"
}""",
    requirements=[
        "Return decimal strings for all money fields, including zero values such as \"0.00\".",
        "Calculate totals server-side from current prices and validated selected options.",
        "Use one currency field consistently; the current frontend expects CHF.",
    ],
)

add_api_section(
    doc, 3, "Add-to-cart deduplication and idempotency", "REQUIRED",
    "Adding the same menu item with the same option configuration currently creates a second line. Identical configurations should be combined, and retries must not duplicate items.",
    """POST /api/app/cart/
Content-Type: multipart/form-data
Idempotency-Key: <unique-operation-id>

data={
  "quantity": 1,
  "menu_item": 364,
  "options": [{"option": 397, "item": 1190}]
}""",
    response="""{
  "id": 42,
  "items": [],
  "item_count": 2,
  "subtotal": "41.00",
  "total_price": "44.50",
  "currency": "CHF"
}""",
    requirements=[
        "Combine lines only when menu item, restaurant, selected options and instructions are identical.",
        "Support Idempotency-Key so network retries return the original result.",
        "Reject products or selected option items that are inactive or unavailable.",
    ],
)

add_api_section(
    doc, 4, "Complete cart-item pricing", "REQUIRED",
    "The cart response contains option identifiers but does not provide reliable resolved option names and prices. Every line should be self-explanatory and carry server-calculated line totals.",
    "GET /api/app/cart/",
    response="""{
  "id": 501,
  "quantity": 2,
  "unit_price": "18.00",
  "options_price": "2.50",
  "line_total": "41.00",
  "menu_item": {
    "id": 364,
    "name": "Classic Burger",
    "description": "Beef burger",
    "price": "18.00",
    "image": "https://cdn.example/menu/364.jpg"
  },
  "options": [{
    "option": 397,
    "option_name": "Choose cheese",
    "item": 1190,
    "item_name": "Gruyère",
    "price": "2.50"
  }]
}""",
)

add_api_section(
    doc, 5, "Clear complete cart", "RECOMMENDED",
    "The frontend currently loops through every line and sends multiple delete requests. A single atomic clear operation is simpler and safer.",
    """DELETE /api/app/cart/items/
Authorization: Bearer <access-token>
X-Session-ID: <guest-session-id>""",
    response="""{
  "id": 42,
  "items": [],
  "item_count": 0,
  "subtotal": "0.00",
  "total_price": "0.00",
  "currency": "CHF"
}""",
    requirements=["Only clear the calling customer or guest session cart."],
)

add_api_section(
    doc, 6, "Guest cart merge after login", "REQUIRED",
    "The guest session is currently cleared during login, which can discard the guest cart. The backend should merge it into the authenticated customer cart.",
    """POST /api/app/cart/merge/
Authorization: Bearer <new-customer-token>
Content-Type: application/json""",
    """{
  "session_id": "guest-session-id"
}""",
    """{
  "id": 42,
  "customer": 18,
  "session_id": null,
  "items": [],
  "subtotal": "52.00",
  "total_price": "55.50",
  "currency": "CHF"
}""",
    [
        "Combine identical configurations.",
        "Resolve or reject carts belonging to different restaurants with a structured conflict response.",
        "Invalidate the guest session only after a successful merge.",
    ],
)

add_api_section(
    doc, 7, "Address update and default management", "REQUIRED",
    "Address list, create and delete APIs exist, but customers cannot edit an address or change the default address cleanly.",
    """PATCH /api/app/customer/addresses/{address_id}/
Authorization: Bearer <access-token>
Content-Type: application/json""",
    """{
  "address": "12 Bahnhofstrasse",
  "postal_code": "8001",
  "city": "Zürich",
  "default": true
}""",
    """{
  "id": 17,
  "address": "12 Bahnhofstrasse",
  "postal_code": "8001",
  "city": "Zürich",
  "default": true
}""",
    [
        "Verify ownership before updating.",
        "If default is true, unset the previous default in the same transaction.",
        "Optionally expose POST /api/app/customer/addresses/{id}/default/ as a dedicated action.",
    ],
)

add_api_section(
    doc, 8, "Address deletion rules", "REQUIRED",
    "The existing delete endpoint needs predictable ownership and default-address behavior.",
    """DELETE /api/app/customer/addresses/{address_id}/
Authorization: Bearer <access-token>""",
    response="""HTTP 204 No Content

// Example validation error
{
  "code": "api.customer_address_last_address",
  "message": "At least one delivery address must remain.",
  "fields": {}
}""",
    requirements=[
        "Never allow deletion of another customer’s address.",
        "When deleting the default, assign another default or return a structured validation error.",
        "Return either 204 consistently or the complete updated list consistently.",
    ],
)

add_api_section(
    doc, 9, "Login role and token refresh", "REQUIRED",
    "The login response does not reliably identify the account role, so the frontend probes profile endpoints. Access tokens also have no implemented refresh flow.",
    """POST /api/token/
Content-Type: application/json

POST /api/token/refresh/
Content-Type: application/json""",
    """// Login
{"username": "customer@example.com", "password": "password"}

// Refresh
{"refresh": "<refresh-token>"}""",
    """{
  "access": "<jwt-access-token>",
  "refresh": "<jwt-refresh-token>",
  "expires_in": 900,
  "user": {
    "id": 18,
    "email": "customer@example.com",
    "firstname": "John",
    "lastname": "Smith",
    "role": "customer"
  }
}""",
    [
        "Return one explicit role: customer, restaurant_owner or admin.",
        "Provide token expiry information and a refresh endpoint.",
        "Prefer a secure, HTTP-only, same-site cookie for the refresh token.",
    ],
)

add_api_section(
    doc, 10, "Customer registration contract", "REQUIRED",
    "Registration should return field-level validation and normalize identity fields. JSON is preferred over a JSON string nested in multipart data when no file is uploaded.",
    """POST /api/app/customer/register/
Content-Type: application/json""",
    """{
  "firstname": "John",
  "lastname": "Smith",
  "birthday": "1995-08-17",
  "address": "12 Bahnhofstrasse",
  "postal_code": "8001",
  "city": "Zürich",
  "email": "john@example.com",
  "phone": "+41791234567",
  "password": "secure-password",
  "password_confirm": "secure-password"
}""",
    """{
  "user": {
    "id": 18,
    "firstname": "John",
    "lastname": "Smith",
    "birthday": "1995-08-17",
    "email": "john@example.com",
    "phone": "+41791234567"
  },
  "access": "<token>",
  "refresh": "<token>"
}

// Validation error
{
  "code": "validation_error",
  "message": "Some details need attention.",
  "fields": {
    "email": ["An account already exists with this email."],
    "phone": ["Enter a valid international phone number."]
  }
}""",
    [
        "Normalize phone numbers to E.164 and emails to a canonical lowercase form.",
        "Enforce uniqueness for normalized email and phone values.",
        "Return stable field-level validation errors.",
    ],
)

add_api_section(
    doc, 11, "Password reset flow", "REQUIRED",
    "The forgot-password page currently simulates success and has no backend integration.",
    """POST /api/app/auth/password-reset/
POST /api/app/auth/password-reset/validate/
POST /api/app/auth/password-reset/confirm/""",
    """// Request reset
{"email": "john@example.com"}

// Validate link
{"uid": "encoded-user-id", "token": "single-use-token"}

// Confirm password
{
  "uid": "encoded-user-id",
  "token": "single-use-token",
  "password": "new-password",
  "password_confirm": "new-password"
}""",
    """// Request response — always generic
{
  "message": "If an account exists, reset instructions have been sent."
}

// Token validation
{"valid": true}""",
    [
        "Use short-lived, cryptographically secure, single-use tokens.",
        "Return the same reset-request message for existing and non-existing accounts.",
        "Rate-limit reset requests by account and client address.",
    ],
)

add_api_section(
    doc, 12, "Restaurant listing, search and filters", "RECOMMENDED",
    "The existing endpoint supports search and category. Additional filters and pagination are needed for a scalable professional listing flow.",
    """GET /api/app/restaurants/
  ?search=pizza
  &category=4
  &city=Zurich
  &open_now=true
  &delivery_available=true
  &pickup_available=true
  &no_waste=false
  &min_rating=4
  &sort=rating
  &page=1
  &page_size=20""",
    response="""{
  "count": 83,
  "next": "/api/app/restaurants/?page=2",
  "previous": null,
  "results": [{
    "id": 7,
    "name": "Restaurant name",
    "delivery_available": true,
    "pickup_available": true,
    "delivery_fee": "3.50",
    "delivery_time": 35,
    "rating": "4.7",
    "reviews": 128,
    "open": true,
    "categories": [],
    "images": []
  }]
}""",
)

add_api_section(
    doc, 13, "Restaurant detail and menu consistency", "REQUIRED",
    "Restaurant detail must consistently expose fulfilment, availability, category, menu option and allergen data so cards do not appear incomplete.",
    "GET /api/app/restaurants/{restaurant_id}/",
    response="""{
  "id": 7,
  "name": "Restaurant",
  "open": true,
  "active": true,
  "delivery_available": true,
  "pickup_available": true,
  "delivery_fee": "3.50",
  "minimum_order": "20.00",
  "estimated_delivery_minutes": 35,
  "foods": [{
    "id": 10,
    "name": "Burgers",
    "image": null,
    "menu_items": [{
      "id": 364,
      "name": "Classic Burger",
      "price": "18.00",
      "image": null,
      "available": true,
      "allergies": [{"id": 2, "name": "Gluten"}],
      "options": [{
        "id": 397,
        "name": "Choose cheese",
        "required": true,
        "multiple": false,
        "minimum_selections": 1,
        "maximum_selections": 1,
        "items": [{
          "id": 1190,
          "name": "Gruyère",
          "price": "2.50",
          "available": true,
          "allergies": []
        }]
      }]
    }]
  }],
  "nowaste_items": []
}""",
    requirements=[
        "Return consistent decimal strings for monetary values.",
        "Include explicit availability for products and option items.",
        "Enforce required, minimum and maximum option selections server-side.",
    ],
)

add_api_section(
    doc, 14, "Image URL reliability", "REQUIRED",
    "Several S3 menu image URLs returned HTTP 403. The frontend now falls back to the Fooddely logo, but the API should not return inaccessible media URLs.",
    "GET /api/app/restaurants/{restaurant_id}/",
    response="""{
  "image": null,
  "image_thumbnail": null
}""",
    requirements=[
        "Return a public CDN URL, a valid signed URL with sufficient expiry, or null.",
        "Do not cache expired signed URLs inside long-lived API responses.",
        "Use thumbnails for listing cards where available.",
    ],
)

add_api_section(
    doc, 15, "Payment intent contract", "REQUIRED",
    "The checkout UI collects fulfilment type, preferred time, delivery notes and notification preference, but the current payment request cannot persist them.",
    """POST /api/app/payment/intent/
Authorization: Bearer <access-token>
X-Session-ID: <guest-session-id>
Content-Type: application/json""",
    """{
  "fulfilment_type": "delivery",
  "delivery_firstname": "John",
  "delivery_lastname": "Smith",
  "delivery_address": "12 Bahnhofstrasse",
  "delivery_postal_code": "8001",
  "delivery_city": "Zürich",
  "delivery_phone": "+41791234567",
  "delivery_email": "john@example.com",
  "preferred_time": "2026-07-29T20:30:00+05:00",
  "delivery_notes": "Ring apartment 4",
  "receive_notifications": false,
  "address_id": 17
}""",
    """{
  "payment_intent_id": "pi_...",
  "payment_intent_client_secret": "pi_..._secret_...",
  "cart_id": 42,
  "subtotal": "52.00",
  "discount_total": "5.00",
  "delivery_fee": "3.50",
  "total_price": "50.50",
  "currency": "CHF"
}""",
    [
        "For pickup, address fields are optional but contact fields remain required.",
        "Validate restaurant availability, delivery area, minimum order and item availability before creating the intent.",
        "Calculate the Stripe amount exclusively on the backend.",
    ],
)

add_api_section(
    doc, 16, "Payment confirmation idempotency", "REQUIRED",
    "Payment confirmation must never create duplicate orders when the frontend retries or Stripe redirects more than once.",
    """POST /api/app/payment/confirm/
Content-Type: application/json""",
    """{
  "payment_intent_id": "pi_..."
}""",
    """{
  "id": 1042,
  "status": "PLACED",
  "payment_intent_id": "pi_...",
  "payment_status": "succeeded",
  "price": "50.50",
  "currency": "CHF",
  "delivery": {},
  "items": []
}""",
    [
        "Verify the PaymentIntent directly with Stripe.",
        "Verify amount and currency against a server-side cart or order snapshot.",
        "Create exactly one order per PaymentIntent; retries return the existing order.",
        "Clear or mark the cart submitted atomically with order creation.",
    ],
)

add_api_section(
    doc, 17, "Checkout validation and quote", "RECOMMENDED",
    "A quote endpoint lets the frontend validate delivery, inventory, restaurant status and discounts before creating a Stripe PaymentIntent.",
    """POST /api/app/cart/quote/
Authorization: Bearer <access-token>
X-Session-ID: <guest-session-id>
Content-Type: application/json""",
    """{
  "fulfilment_type": "delivery",
  "address_id": 17,
  "postal_code": "8001",
  "preferred_time": "2026-07-29T20:30:00+05:00",
  "discount_code": "WELCOME10"
}""",
    """{
  "valid": true,
  "errors": [],
  "subtotal": "52.00",
  "options_total": "4.50",
  "discount_total": "5.00",
  "delivery_fee": "3.50",
  "total_price": "55.00",
  "currency": "CHF",
  "minimum_order_met": true,
  "restaurant_open": true,
  "items_available": true
}""",
)

add_api_section(
    doc, 18, "Standard API error envelope", "REQUIRED",
    "All endpoints should return a stable machine-readable error contract so the frontend can show useful field and workflow messages.",
    "Applies to all /api/app/* endpoints",
    response="""{
  "code": "api.cart_item_unavailable",
  "message": "This item is no longer available.",
  "fields": {
    "quantity": ["Only 2 portions are currently available."]
  }
}""",
    requirements=[
        "400 — malformed request.",
        "401 — authentication missing or expired.",
        "403 — authenticated but not permitted.",
        "404 — resource not found.",
        "409 — cart conflict, duplicate operation or unavailable inventory.",
        "422 — field validation failure.",
        "429 — rate limited.",
    ],
)

doc.add_heading("Backend acceptance checklist", level=1)
checklist = [
    "Cart quantity can be changed without delete-and-recreate behavior.",
    "Every cart response includes non-null totals and currency.",
    "Identical add-to-cart retries do not create duplicate lines.",
    "Guest cart survives login and merges predictably.",
    "Addresses can be edited and one default is maintained transactionally.",
    "Login returns an explicit role and tokens can be refreshed.",
    "Password reset works end-to-end with single-use tokens.",
    "Restaurant and menu media URLs are browser-accessible or null.",
    "Checkout persists delivery/pickup, time and notes.",
    "Stripe amount is generated and validated only by the backend.",
    "Payment confirmation is idempotent and cannot create duplicate orders.",
    "Errors use one consistent response envelope.",
]
for item in checklist:
    add_bullet(doc, f"☐ {item}")

doc.add_heading("Implementation order", level=1)
for idx, text in enumerate([
    "Cart mutation and complete pricing response.",
    "Payment intent and idempotent confirmation.",
    "Guest-cart merge and authentication refresh.",
    "Password reset and address update/default management.",
    "Image reliability and restaurant/menu payload normalization.",
    "Pagination, expanded filters, clear-cart and quote endpoint.",
], start=1):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    format_run(p.add_run(text), size=10.25)

doc.core_properties.title = "Fooddely Backend API Requirements"
doc.core_properties.subject = "Required backend changes and new API endpoints"
doc.core_properties.author = "Fooddely Frontend Team"
doc.core_properties.keywords = "Fooddely, API, backend, cart, authentication, checkout, Stripe"

doc.save(OUTPUT)
print(OUTPUT)
